import io
from pathlib import Path

from openpyxl import load_workbook
import pytest

from backend.services.generator import generate_template
from backend.services.processor import process_file_data
from backend.parsers.txt_parser import parse_txt_content

pytestmark = pytest.mark.compliance


def test_template_flow_generates_compliant_txt():
    template_buffer = generate_template()
    wb = load_workbook(template_buffer)

    ws_header = wb["Cabecalho"]
    header_vals = [
        "123456",
        "22/01/2026",
        "23/01/2026",
        "3101",
        "PO-999",
        "0.00",
        "IMP-001",
        "TRA-002",
        "TEST EXPORTER LTD",
    ]
    for col_idx, val in enumerate(header_vals, start=1):
        ws_header.cell(row=2, column=col_idx, value=val)

    ws_items = wb["Itens"]
    item_rows = [
        ["85044090", "Desc Test", "PROD-001", "UN", "3101"],
        ["85176272", "Desc Test 2", "PROD-002", "UN", "3101"],
        ["85367000", "Desc Test 3", "PROD-003", "UN", "3101"],
        ["90283031", "Desc Test 4", "PROD-004", "UN", "3101"],
    ]
    for row_idx, row_vals in enumerate(item_rows, start=2):
        for col_idx, val in enumerate(row_vals, start=1):
            ws_items.cell(row=row_idx, column=col_idx, value=val)

    input_xlsx = io.BytesIO()
    wb.save(input_xlsx)
    input_xlsx.seek(0)

    base_dir = Path(__file__).resolve().parent.parent.parent
    xml_path = base_dir / "format_guidelines" / "Eletra.xml"
    xml_content = xml_path.read_bytes()

    txt_output = process_file_data(xml_content, input_xlsx.getvalue())
    assert txt_output.endswith("\r\n")

    lines = txt_output.splitlines()
    assert lines
    assert len(lines[0]) == 380
    assert all(len(line) == 717 for line in lines[1:])

    header_line = lines[0]
    assert header_line[5:14] == "000123456"
    assert header_line[80:88] == "20260122"
    assert header_line[297:357].strip()

    parsed_header, parsed_items = parse_txt_content(txt_output.encode("utf-8"))
    assert parsed_header["numero_nota"] == "123456"
    assert parsed_items
