import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

def generate_template() -> io.BytesIO:
    """
    Generates a detailed Excel template for the user to fill.
    Includes instructions, examples, and validation guidance.
    """
    output = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # --- Styling Definitions ---
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    example_font = Font(italic=True, color="666666")
    title_font = Font(bold=True, size=14)
    h2_font = Font(bold=True, size=12)
    
    # --- Sheet 1: Instructions ---
    ws_instructions = wb.create_sheet(title="Instrucoes", index=0)
    
    ws_instructions.merge_cells('A1:D1')
    ws_instructions['A1'].value = "Guia de Preenchimento do Modelo"
    ws_instructions['A1'].font = title_font
    
    ws_instructions['A3'].value = "Objetivo do Modelo:"
    ws_instructions['A3'].font = h2_font
    ws_instructions['A4'].value = "Este modelo serve para complementar os dados da Declaração de Importação (DI) XML com informações internas do seu sistema (códigos de produto, CFOPs) e para definir dados da nota fiscal que não existem no XML."

    ws_instructions['A6'].value = "Guia Rápido:"
    ws_instructions['A6'].font = h2_font
    ws_instructions['A7'].value = "1. Preencha a aba 'Cabecalho' com os dados da sua nota fiscal (Número e Datas)."
    ws_instructions['A8'].value = "2. Na aba 'Itens', mapeie cada NCM do seu XML para o código de produto interno e CFOP correspondente."
    ws_instructions['A9'].value = "3. Salve o arquivo e faça o upload na aplicação junto com o XML da DI."
    ws_instructions['A10'].value = "4. Importante: A ordem das colunas deve ser mantida. Não adicione, remova ou reordene."
    ws_instructions['A11'].value = "5. A ordem das linhas na aba 'Itens' não importa."

    ws_instructions['A13'].value = "Histórico de Versões:"
    ws_instructions['A13'].font = h2_font
    ws_instructions.cell(row=14, column=1).value = "Versão"
    ws_instructions.cell(row=14, column=2).value = "Data"
    ws_instructions.cell(row=14, column=3).value = "Alteração"
    ws_instructions.cell(row=15, column=1).value = "1.0"
    ws_instructions.cell(row=15, column=2).value = "22/01/2026"
    ws_instructions.cell(row=15, column=3).value = "Versão inicial."
    ws_instructions.cell(row=16, column=1).value = "1.1"
    ws_instructions.cell(row=16, column=2).value = "23/01/2026"
    ws_instructions.cell(row=16, column=3).value = "Adicionado guia de preenchimento e exemplos."
    ws_instructions.cell(row=17, column=1).value = "1.2"
    ws_instructions.cell(row=17, column=2).value = "23/01/2026"
    ws_instructions.cell(row=17, column=3).value = "Adicionada nota sobre a ordem das colunas e linhas."
    ws_instructions.cell(row=18, column=1).value = "2.0"
    ws_instructions.cell(row=18, column=2).value = "23/01/2026"
    ws_instructions.cell(row=18, column=3).value = "Simplificado template. Frete, Seguro e Exportador agora são lidos do XML."

    # --- Sheet 2: Cabecalho ---
    ws_header = wb.create_sheet(title="Cabecalho", index=1)
    
    headers_cabecalho = {
        "A1": ("Numero da Nota (Obrigatorio)", "Número da nota fiscal de entrada que será gerada. Ex: 12345"),
        "B1": ("Data Emissao (DD/MM/AAAA)", "Data de emissão da nota."),
        "C1": ("Data Entrada (DD/MM/AAAA)", "Data de entrada das mercadorias."),
        "D1": ("CFOP Predominante (Obrigatorio)", "CFOP principal da nota (Ex: 3101). Irá no cabeçalho."),
        "E1": ("Numero do PO (Opcional)", "Número do Pedido de Compra (Campo 44 do Layout)."),
        "F1": ("Perc Reducao ICMS (Opcional)", "Percentual de redução da base de ICMS (Campo 45)."),
        "G1": ("Codigo Importador (Opcional)", "Código do Importador no sistema destino."),
        "H1": ("Codigo Transportador (Opcional)", "Código do Transportador no sistema destino."),
        "I1": ("Codigo Exportador (Opcional)", "Código do Exportador no sistema destino."),
    }
    
    for cell, (title, comment_text) in headers_cabecalho.items():
        ws_header[cell].value = title
        ws_header[cell].comment = Comment(comment_text, "Sistema")
        ws_header[cell].font = header_font
        ws_header[cell].fill = header_fill
        ws_header[cell].alignment = Alignment(horizontal="center", vertical="center")
        ws_header.column_dimensions[cell[0]].width = 25

    # Example Row
    example_header = ["12345", "23/01/2026", "24/01/2026", "3101", "PO-999", "0.00", "IMP-001", "TRA-002", "EXP-003"]
    ws_header.append(example_header)
    for cell in ws_header[2]:
        cell.font = example_font


    # --- Sheet 3: Itens ---
    ws_items = wb.create_sheet(title="Itens", index=2)
    
    headers_items = {
        "A1": ("Codigo NCM (XML)", "O código NCM de 8 dígitos presente no XML. Este é o campo chave para o mapeamento."),
        "B1": ("Descricao (Informativo)", "Descrição do produto. Este campo é apenas para referência e não é usado pelo sistema."),
        "C1": ("Codigo Produto Interno", "O código do produto no seu sistema ERP. (Obrigatório)"),
        "D1": ("Unidade Saida", "Unidade de medida que deve constar na nota (ex: UN, CX, KG). Se vazio, usará a unidade do XML."),
        "E1": ("CFOP", "CFOP de entrada da mercadoria. (Obrigatório)")
    }
    
    for cell, (title, comment_text) in headers_items.items():
        ws_items[cell].value = title
        ws_items[cell].comment = Comment(comment_text, "Sistema")
        ws_items[cell].font = header_font
        ws_items[cell].fill = header_fill
        ws_items[cell].alignment = Alignment(horizontal="center", vertical="center")

    widths = [20, 40, 30, 25, 20]
    for i, w in enumerate(widths):
        ws_items.column_dimensions[get_column_letter(i+1)].width = w

    # Example Row
    example_items = ["85171231", "Smartphone Modelo X", "PROD-00123", "UN", "3101"]
    ws_items.append(example_items)
    for cell in ws_items[2]:
        cell.font = example_font
        
    # Finalize
    wb.active = ws_instructions
    wb.save(output)
    output.seek(0)
    return output
